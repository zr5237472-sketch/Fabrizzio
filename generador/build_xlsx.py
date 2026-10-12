import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import LineChart, Reference
from data import *

import os
OUT=os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),"Seguimiento_12_semanas.xlsx")
wb=Workbook()
GREEN="2F6B4F"; LIGHT="E8F1EC"; INPUT="DDEBF7"; GREY="F2F2F2"
hfont=Font(bold=True,color="FFFFFF"); hfill=PatternFill("solid",fgColor=GREEN)
thin=Side(style="thin",color="BFBFBF"); box=Border(left=thin,right=thin,top=thin,bottom=thin)
blue=Font(color="0000FF"); bold=Font(bold=True)
def hdr(ws,row,cols,start=1):
    for i,t in enumerate(cols):
        c=ws.cell(row,start+i,t); c.font=hfont; c.fill=hfill; c.border=box
        c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    ws.row_dimensions[row].height=34
def title(ws,t,sub=None):
    ws["A1"]=t; ws["A1"].font=Font(bold=True,size=15,color=GREEN)
    if sub: ws["A2"]=sub; ws["A2"].font=Font(italic=True,color="595959")
def widths(ws,ws_):
    for i,w in enumerate(ws_,1): ws.column_dimensions[L(i)].width=w
def inp(c):
    c.font=blue; c.fill=PatternFill("solid",fgColor=INPUT); c.border=box

# ================= LEEME =================
ws=wb.active; ws.title="Leeme"
title(ws,"Seguimiento nutricional · 12 semanas","Cliente: mujer adulta · 165 cm · 64,5 kg → meta 61 kg · entrena 2–3 días/semana")
rows=[
("Cómo usar este libro",""),
("1. Perfil","Datos y supuestos (celdas azules = editables). Cambia edad, factor de actividad o gramos por medida casera y todo se recalcula."),
("2. Plan_12sem","Peso objetivo, calorías y macros por semana (fórmulas Mifflin-St Jeor), palancas de ajuste y pérdida esperada."),
("3. Alimentos","Base TPCA (INS, Perú) por 100 g. Es la tabla que alimenta todos los cálculos."),
("4. Menu_S1 / Resumen_S1","Menú planificado de la semana 1 con gramos, kcal y macros por día."),
("5. Registro","CADA FIN DE SEMANA: llenar semana, día, comida, alimento (lista), cantidad y unidad según la foto. Calcula solo."),
("6. Dia_Detalle","Elige la semana en C3 y ve día por día cuánto cumplió, cuánto falta o sobra."),
("7. Control_Semanal","Resumen de las 12 semanas: cumplimiento, faltante de kcal y proteína."),
("8. Peso","Registrar peso cada semana y rendimiento en gym (1–5). Incluye gráfico objetivo vs real."),
("9. Reglas","Palancas de ajuste y criterios de decisión para proteger el rendimiento."),
("",""),
("Supuestos importantes",""),
("Edad","No fue indicada: se asume 25 años (afecta ±50 kcal). Corregir en Perfil."),
("Actividad","Factor 1,35 = estudiante + gym 2–3 días/semana."),
("Scoop de avena","Se asume 25 g de hojuela cruda por scoop (verificar con balanza una vez)."),
("Rebanada de pan de molde","Se asume 25 g por rebanada (verificar con el pan que usa)."),
("Cucharón de arroz","70 g de arroz cocido por cucharón (3 cucharones ≈ 210 g)."),
("Papa grande","150 g pelada y sancochada (2 papas ≈ 300 g)."),
("Tomatodo","800 ml: 3 tomatodos ≈ 2,4 L."),
("Estimación por foto","Es aproximada (±10–15 %). Sirve para tendencia; el peso semanal es el control real."),
]
for i,(a,b) in enumerate(rows,4):
    ws.cell(i,1,a).font=bold if b=="" or a[:1].isdigit() else Font(bold=True)
    ws.cell(i,2,b).alignment=Alignment(wrap_text=True,vertical="top")
widths(ws,[28,110])
for r in (4,15): ws.cell(r,1).font=Font(bold=True,size=12,color=GREEN)

# ================= PERFIL =================
pf=wb.create_sheet("Perfil"); title(pf,"Perfil y supuestos","Celdas azules = editables")
items=[("Sexo","Mujer"),("Edad (años) — supuesto",PERFIL["edad"]),("Talla (cm)",PERFIL["talla_cm"]),
("Peso inicial (kg)",PERFIL["peso0"]),("Peso meta semana 12 (kg)",PERFIL["peso_meta"]),
("Factor de actividad (PAL)",PERFIL["pal"]),("Proteína (g/kg)",PERFIL["g_prot_kg"]),
("Grasa (g/kg)",0.6),("Piso de calorías (kcal/día)",PERFIL["piso_kcal"]),
("Inicio semana 1 (lunes)",INICIO),("kcal por kg de peso",7700),("Tomatodo (ml)",800)]
for i,(a,b) in enumerate(items,4):
    pf.cell(i,2,a); c=pf.cell(i,3,b); inp(c)
    if isinstance(b,dt.date): c.number_format="dd/mm/yyyy"
# C4 sexo, C5 edad, C6 talla, C7 peso0, C8 meta, C9 PAL, C10 prot, C11 grasa, C12 piso, C13 inicio, C14 kcal/kg, C15 tomatodo
pf["B18"]="Unidad"; pf["C18"]="Gramos"; hdr(pf,18,["Unidad (medida casera)","Gramos por unidad"],2)
meds=[("rebanada",MEDIDAS["rebanada"]),("scoop",MEDIDAS["scoop"]),("cucharón",MEDIDAS["cucharon"]),
      ("papa grande",MEDIDAS["papa"]),("cdta",5),("g",1)]
for i,(a,b) in enumerate(meds,19):
    pf.cell(i,2,a); inp(pf.cell(i,3,b))
MED_RNG="Perfil!$B$19:$C$24"
pf["E19"]="Rebanada de pan de molde"; pf["E20"]="Scoop de avena (hojuela cruda)"; pf["E21"]="Cucharón de arroz cocido"
pf["E22"]="Papa grande pelada y sancochada"; pf["E23"]="Cucharadita de aceite"; pf["E24"]="Gramos (directo)"
widths(pf,[3,34,16,3,34])

# ================= ALIMENTOS =================
al=wb.create_sheet("Alimentos"); title(al,"Base de alimentos (por 100 g)","Fuente: Tablas Peruanas de Composición de Alimentos (INS/CENAN). Valores editables.")
hdr(al,4,["Clave","Alimento","Código TPCA","Estado","kcal","Proteína (g)","Grasa (g)","Carbohidrato (g)","Categoría"])
cat={"pan_molde":"Desayuno/merienda","pan_frances":"Desayuno/merienda","avena":"Desayuno/merienda","quinua_hoj":"Desayuno/merienda",
"arroz":"Carbohidrato","papa":"Carbohidrato","papa_amar":"Carbohidrato","camote":"Carbohidrato","fideo":"Carbohidrato","quinua_cocida":"Carbohidrato","quinua_coc":"Carbohidrato",
"queso":"Grasa/lácteo","palta":"Grasa/lácteo","aceite":"Grasa/lácteo","lenteja":"Proteína vegetal","garbanzo":"Proteína vegetal"}
for i,(k,v) in enumerate(F.items(),5):
    n,kc,p,g,c,code,est=v
    cg=cat.get(k) or ("Proteína" if k in("pollo_pech","pollo_pulpa","res_pulpa","tilapia","trucha","bonito","merluza","atun_agua","huevo","clara") else "Verdura")
    for j,val in enumerate([k,n,code,est,kc,p,g,c,cg],1):
        cell=al.cell(i,j,val); cell.border=box
        if j in (5,6,7,8): cell.font=blue
NA=4+len(F)
widths(al,[13,44,12,14,9,12,10,15,20]); al.freeze_panes="C5"
AN=f"Alimentos!$B$5:$B${NA}"
def look(col,key): return f"INDEX(Alimentos!${col}$5:${col}${NA},MATCH({key},{AN},0))"

# ================= MENU_S1 =================
mn=wb.create_sheet("Menu_S1"); title(mn,"Menú semana 1 · 12 al 18 de octubre","Cantidad × unidad = gramos. Valores nutricionales vía base TPCA.")
hdr(mn,4,["Día","Comida","Alimento","Cantidad","Unidad","Gramos","kcal","Proteína (g)","Grasa (g)","Carbohidrato (g)"])
def tidy(key,gr):
    if key=="pan_molde": return gr/MEDIDAS["rebanada"],"rebanada"
    if key=="avena": return gr/MEDIDAS["scoop"],"scoop"
    if key=="arroz": return gr/MEDIDAS["cucharon"],"cucharón"
    if key=="papa": return gr/MEDIDAS["papa"],"papa grande"
    if key=="aceite": return gr/5,"cdta"
    return gr,"g"
r=5; NOMB={"des":"Desayuno","alm":"Almuerzo","mer":"Merienda","cen":"Cena"}
for dn,d in zip(DIAS,MENU):
    ci=comida_items(d)
    for cm in ("des","alm","mer","cen"):
        for k,gr in ci[cm].items():
            q,u=tidy(k,gr)
            mn.cell(r,1,dn); mn.cell(r,2,NOMB[cm]); mn.cell(r,3,F[k][0])
            c=mn.cell(r,4,round(q,2)); inp(c); mn.cell(r,5,u).font=blue
            mn.cell(r,6,f"=D{r}*IFERROR(VLOOKUP(E{r},{MED_RNG},2,0),1)").number_format="0"
            for j,col in zip(range(7,11),"EFGH"):
                mn.cell(r,j,f"=F{r}/100*{look(col,f'$C{r}')}").number_format="0.0"
            for j in range(1,11): mn.cell(r,j).border=box
            r+=1
MN_LAST=r-1
widths(mn,[12,12,44,10,13,9,9,12,10,15]); mn.freeze_panes="A5"
mn.auto_filter.ref=f"A4:J{MN_LAST}"
# ================= RESUMEN_S1 =================
rs=wb.create_sheet("Resumen_S1"); title(rs,"Resumen semana 1: plan vs meta","Meta diaria tomada de Plan_12sem (fila semana 1)")
hdr(rs,4,["Día","Desayuno","Almuerzo","Merienda","Cena","Total kcal","Proteína (g)","Grasa (g)","Carbohidrato (g)","Meta kcal/día","Diferencia vs meta","Meta proteína (g)"])
for i,dn in enumerate(DIAS,5):
    rs.cell(i,1,dn)
    for j,cm in enumerate(("Desayuno","Almuerzo","Merienda","Cena"),2):
        rs.cell(i,j,f'=SUMIFS(Menu_S1!$G$5:$G${MN_LAST},Menu_S1!$A$5:$A${MN_LAST},$A{i},Menu_S1!$B$5:$B${MN_LAST},"{cm}")').number_format="0"
    rs.cell(i,6,f"=SUM(B{i}:E{i})").number_format="0"
    for j,col in zip((7,8,9),"HIJ"):
        rs.cell(i,j,f'=SUMIFS(Menu_S1!${col}$5:${col}${MN_LAST},Menu_S1!$A$5:$A${MN_LAST},$A{i})').number_format="0"
    rs.cell(i,10,"=Plan_12sem!$J$5").number_format="0"
    rs.cell(i,11,f"=F{i}-J{i}").number_format="+0;-0;0"
    rs.cell(i,12,"=Plan_12sem!$L$5")
    for j in range(1,13): rs.cell(i,j).border=box
rs.cell(12,1,"Promedio").font=bold
for j in range(2,13):
    c=rs.cell(12,j,f"=AVERAGE({L(j)}5:{L(j)}11)"); c.number_format="0"; c.font=bold; c.border=box
rs.cell(14,1,"Proteína por kg (promedio)").font=bold; rs.cell(14,3,"=G12/Perfil!C7").number_format="0.00"
rs.cell(15,1,"Pérdida esperada con este plan (kg/sem)").font=bold
rs.cell(15,3,"=(Plan_12sem!H5-F12)*7/Perfil!C14").number_format="0.00"
rs.cell(16,1,"Nota: los días con queso (miércoles y sábado) suben ≈ +250 kcal; por eso el promedio queda sobre la meta de la semana 1.").font=Font(italic=True)
widths(rs,[26,11,11,11,11,11,12,10,15,13,14,14])
BASE="Resumen_S1!$F$12"

# ================= PLAN_12SEM =================
pl=wb.create_sheet("Plan_12sem",1); title(pl,"Plan de 12 semanas","64,5 kg → 61 kg · descenso progresivo (0,20 → 0,25 → 0,30 kg/sem) para cuidar el rendimiento en gym")
cols=["Sem","Inicio","Fin","Peso objetivo fin de sem (kg)","Peso inicio de sem (kg)","Cambio esperado (kg)",
"BMR (kcal)","Gasto total GET (kcal)","Déficit (kcal/día)","META kcal/día","Meta kcal/semana","Proteína (g/día)",
"Grasa (g/día)","Carbo (g/día)","Agua (L/día)","Tomatodos","L1 queso 100 g","L2 cena carbo 1/4","L3 almuerzo carbo 1/4","L4 merienda avena 1,5","L5 sin aceite",
"Recorte vs plan S1 (kcal/día)","kcal estimadas del plan","Brecha vs meta (kcal)","Pérdida esperada (kg/sem)","Notas"]
hdr(pl,4,cols); pl.row_dimensions[4].height=62
wts=[64.3,64.05,63.75,63.45,63.15,62.85,62.55,62.25,61.95,61.65,61.35,61.0]
flags={1:[0,0,0,0,0],2:[1,1,0,0,0],3:[1,1,1,0,0],4:[1,1,1,1,0]}
for w in range(5,13): flags[w]=[1,1,1,1,1]
notas={1:"Plan base del cliente. Tomar fotos fin de semana.",2:"Primeros recortes suaves.",3:"",4:"",5:"",6:"Pesar 2 semanas seguidas antes de recortar más.",
 7:"Si el gym baja: +150 kcal pre-entreno (merienda) por 1 semana.",8:"",9:"",10:"",11:"",12:"Cierre: ver Reglas para mantenimiento."}
for w in range(1,13):
    r=4+w
    pl.cell(r,1,w)
    pl.cell(r,2,f"=Perfil!$C$13+(A{r}-1)*7").number_format="dd/mm"
    pl.cell(r,3,f"=B{r}+6").number_format="dd/mm"
    c=pl.cell(r,4,wts[w-1]); inp(c); c.number_format="0.00"
    pl.cell(r,5,"=Perfil!$C$7" if w==1 else f"=D{r-1}").number_format="0.00"
    pl.cell(r,6,f"=E{r}-D{r}").number_format="0.00"
    pl.cell(r,7,f"=10*E{r}+6.25*Perfil!$C$6-5*Perfil!$C$5-161").number_format="0"
    pl.cell(r,8,f"=G{r}*Perfil!$C$9").number_format="0"
    pl.cell(r,9,f"=F{r}*Perfil!$C$14/7").number_format="0"
    c=pl.cell(r,10,f"=MAX(Perfil!$C$12,ROUND(H{r}-I{r},-1))"); c.number_format="0"; c.font=bold; c.fill=PatternFill("solid",fgColor=LIGHT)
    pl.cell(r,11,f"=J{r}*7").number_format="#,##0"
    pl.cell(r,12,f"=ROUND(E{r}*Perfil!$C$10,0)")
    pl.cell(r,13,f"=ROUND(E{r}*Perfil!$C$11,0)")
    pl.cell(r,14,f"=ROUND((J{r}-4*L{r}-9*M{r})/4,0)")
    pl.cell(r,15,f"=ROUND(E{r}*0.037,1)").number_format="0.0"
    pl.cell(r,16,f"=ROUND(O{r}*1000/Perfil!$C$15,1)").number_format="0.0"
    for j,f in enumerate(flags[w],17): inp(pl.cell(r,j,f))
    pl.cell(r,22,f"=Q{r}*Reglas!$C$6+R{r}*Reglas!$C$7+S{r}*Reglas!$C$8+T{r}*Reglas!$C$9+U{r}*Reglas!$C$10").number_format="0"
    pl.cell(r,23,f"={BASE}-V{r}").number_format="0"
    pl.cell(r,24,f"=W{r}-J{r}").number_format="+0;-0;0"
    pl.cell(r,25,f"=(H{r}-W{r})*7/Perfil!$C$14").number_format="0.00"
    pl.cell(r,26,notas[w])
    for j in range(1,27): pl.cell(r,j).border=box
pl.cell(18,1,"Lectura: META kcal es el objetivo fisiológico; 'kcal estimadas del plan' es lo que da el menú con las palancas activas. Brecha positiva = el plan aún está sobre la meta.").font=Font(italic=True)
pl.cell(19,1,"Peso objetivo de la semana 12: 61,0 kg. Las variaciones de ±0,5 kg por agua/ciclo son normales: decidir con el promedio de 2 semanas.").font=Font(italic=True)
widths(pl,[6,9,9,13,13,11,9,11,10,11,11,11,10,10,9,10,10,10,10,10,9,12,12,11,12,48]); pl.freeze_panes="B5"
pl.conditional_formatting.add("X5:X16",CellIsRule(operator="greaterThan",formula=["60"],fill=PatternFill("solid",bgColor="F8CBAD")))

# ================= REGLAS =================
rg=wb.create_sheet("Reglas"); title(rg,"Palancas y reglas de decisión","Orden de aplicación: de menor a mayor impacto sobre el rendimiento")
hdr(rg,5,["Palanca","Qué se hace","kcal recortadas (promedio/día)","Cuándo"])
lev=[("L1","Queso fresco 150 g → 100 g (2 días/sem)",38,"Semana 2 en adelante"),
("L2","Cena: carbohidrato de 1/3 a 1/4 de plato",55,"Semana 2"),
("L3","Almuerzo: carbohidrato de 1/3 a 1/4 de plato (mantener en días de gym)",60,"Semana 3"),
("L4","Merienda: avena 2 → 1,5 scoops (solo días SIN gym)",42,"Semana 4"),
("L5","Sin aceite de oliva (cocinar en plancha antiadherente)",44,"Semana 5")]
for i,(a,b,c,d) in enumerate(lev,6):
    rg.cell(i,1,a); rg.cell(i,2,b); x=rg.cell(i,3,c); inp(x); rg.cell(i,4,d)
    for j in range(1,5): rg.cell(i,j).border=box
rg["A13"]="Reglas de decisión semanales"; rg["A13"].font=Font(bold=True,size=12,color=GREEN)
reglas=[
"Pesar siempre en las mismas condiciones (mañana, en ayunas, tras ir al baño). Usar el promedio de 2 semanas.",
"Peso real ≤ objetivo + 0,2 kg → sin cambios.",
"Peso real > objetivo + 0,3 kg durante 2 semanas → activar la siguiente palanca no activa (L1 a L5).",
"Baja > 0,6 kg en una semana o cae el rendimiento (fuerza, repeticiones) 2 semanas seguidas → +150 kcal de carbohidrato pre-entreno por 1 semana.",
"Proteína no se recorta: mantener ≈ 1,6 g/kg para proteger masa muscular.",
"Hidratación: 3 tomatodos de agua mínimo, más en días de gym.",
"Semana 12: pasar a mantenimiento (≈ GET) por 2 semanas antes de decidir si se baja más.",
"Cumplimiento por foto: 'Cumple' si el promedio diario está dentro de ±5 % de la meta; 'Falta/Exceso' fuera de ese rango.",
]
for i,t in enumerate(reglas,14): rg.cell(i,1,f"{i-13}."); rg.cell(i,2,t).alignment=Alignment(wrap_text=True)
widths(rg,[9,78,20,28])

# ================= REGISTRO =================
rgst=wb.create_sheet("Registro"); title(rgst,"Registro semanal por foto","Llenar cada fin de semana: una fila por alimento visto en la foto. Unidad 'g' para pesos directos.")
rgst["A3"]="Tip: si solo se ve el plato, estimar con la regla 1/3 – 1/3 – 1/3 y la palma (150 g crudo)."; rgst["A3"].font=Font(italic=True)
hdr(rgst,5,["Semana","Día","Comida","Alimento","Cantidad","Unidad","Gramos","kcal","Proteína (g)","Grasa (g)","Carbo (g)","Fecha"])
R0,R1=6,1005
dv_d=DataValidation(type="list",formula1='"Lunes,Martes,Miércoles,Jueves,Viernes,Sábado,Domingo"',allow_blank=True)
dv_c=DataValidation(type="list",formula1='"Desayuno,Almuerzo,Merienda,Cena,Extra"',allow_blank=True)
dv_a=DataValidation(type="list",formula1=f"={AN}",allow_blank=True)
dv_u=DataValidation(type="list",formula1='"g,rebanada,scoop,cucharón,papa grande,cdta"',allow_blank=True)
dv_s=DataValidation(type="whole",operator="between",formula1="1",formula2="12",allow_blank=True)
for dv in (dv_d,dv_c,dv_a,dv_u,dv_s): rgst.add_data_validation(dv)
dv_s.add(f"A{R0}:A{R1}"); dv_d.add(f"B{R0}:B{R1}"); dv_c.add(f"C{R0}:C{R1}"); dv_a.add(f"D{R0}:D{R1}"); dv_u.add(f"F{R0}:F{R1}")
for r in range(R0,R1+1):
    rgst.cell(r,7,f'=IF(OR(D{r}="",E{r}=""),"",E{r}*IFERROR(VLOOKUP(F{r},{MED_RNG},2,0),1))')
    for j,col in zip(range(8,12),"EFGH"):
        rgst.cell(r,j,f'=IF(G{r}="","",G{r}/100*{look(col,f"$D{r}")})')
    rgst.cell(r,12,f'=IF(OR(A{r}="",B{r}=""),"",Perfil!$C$13+(A{r}-1)*7+MATCH(B{r},{{"Lunes","Martes","Miércoles","Jueves","Viernes","Sábado","Domingo"}},0)-1)').number_format="dd/mm"
    for j in (7,8,9,10,11): rgst.cell(r,j).number_format="0"
widths(rgst,[9,12,12,44,10,13,9,9,12,10,10,9]); rgst.freeze_panes="A6"
def RG(col): return f"Registro!${col}${R0}:${col}${R1}"

# ================= DIA_DETALLE =================
dd=wb.create_sheet("Dia_Detalle"); title(dd,"Detalle por día","Elige la semana en C3")
dd["B3"]="Semana (1–12):"; dd["B3"].font=bold; dd["C3"]=1; inp(dd["C3"])
dvw=DataValidation(type="whole",operator="between",formula1="1",formula2="12"); dd.add_data_validation(dvw); dvw.add("C3")
dd["E3"]="Meta kcal/día:"; dd["F3"]="=INDEX(Plan_12sem!$J$5:$J$16,$C$3)"; dd["F3"].font=bold
dd["H3"]="Meta proteína:"; dd["I3"]="=INDEX(Plan_12sem!$L$5:$L$16,$C$3)"; dd["I3"].font=bold
hdr(dd,5,["Día","kcal reales","Meta kcal","Dif. (real − meta)","% cumplimiento","Estado","Proteína real (g)","Meta proteína","Falta proteína (g)","Grasa (g)","Carbo (g)"])
for i,dn in enumerate(DIAS,6):
    dd.cell(i,1,dn)
    dd.cell(i,2,f'=SUMIFS({RG("H")},{RG("A")},$C$3,{RG("B")},$A{i})').number_format="0"
    dd.cell(i,3,"=$F$3")
    dd.cell(i,4,f"=IF(B{i}=0,\"\",B{i}-C{i})").number_format="+0;-0;0"
    dd.cell(i,5,f"=IF(B{i}=0,\"\",B{i}/C{i})").number_format="0%"
    dd.cell(i,6,f'=IF(B{i}=0,"Sin registro",IF(ABS(E{i}-1)<=0.05,"Cumple",IF(E{i}>1.05,"Exceso","Falta")))')
    dd.cell(i,7,f'=SUMIFS({RG("I")},{RG("A")},$C$3,{RG("B")},$A{i})').number_format="0"
    dd.cell(i,8,"=$I$3")
    dd.cell(i,9,f'=IF(G{i}=0,"",MAX(0,H{i}-G{i}))').number_format="0"
    dd.cell(i,10,f'=SUMIFS({RG("J")},{RG("A")},$C$3,{RG("B")},$A{i})').number_format="0"
    dd.cell(i,11,f'=SUMIFS({RG("K")},{RG("A")},$C$3,{RG("B")},$A{i})').number_format="0"
    for j in range(1,12): dd.cell(i,j).border=box
dd.cell(13,1,"Semana").font=bold
dd.cell(13,2,"=SUM(B6:B12)").number_format="0"; dd.cell(13,3,'=C6*COUNTIF(B6:B12,">0")').number_format="0"
dd.cell(13,4,'=IF(C13=0,"",B13-C13)').number_format="+0;-0;0"
dd.cell(13,5,'=IF(C13=0,"",B13/C13)').number_format="0%"
dd.cell(13,7,"=SUM(G6:G12)").number_format="0"; dd.cell(13,8,'=H6*COUNTIF(G6:G12,">0")')
dd.cell(13,9,'=IF(H13=0,"",MAX(0,H13-G13))').number_format="0"
for j in range(1,12): dd.cell(13,j).font=bold; dd.cell(13,j).border=box
dd["A15"]="Nota: la fila Semana compara solo los días registrados. Si faltan días, completa el Registro."; dd["A15"].font=Font(italic=True)
widths(dd,[14,12,11,16,14,14,15,13,15,11,11])
for rng in ("F6:F12",):
    dd.conditional_formatting.add(rng,CellIsRule(operator="equal",formula=['"Cumple"'],fill=PatternFill("solid",bgColor="C6EFCE")))
    dd.conditional_formatting.add(rng,CellIsRule(operator="equal",formula=['"Exceso"'],fill=PatternFill("solid",bgColor="F8CBAD")))
    dd.conditional_formatting.add(rng,CellIsRule(operator="equal",formula=['"Falta"'],fill=PatternFill("solid",bgColor="FFEB9C")))

# ================= CONTROL_SEMANAL =================
cs=wb.create_sheet("Control_Semanal"); title(cs,"Control semanal (S1–S12)","Promedios sobre los días registrados")
hdr(cs,4,["Sem","Inicio","Meta kcal/día","Días registrados","kcal registradas","Promedio real kcal/día","Dif. vs meta/día","% cumplimiento","Estado",
"Proteína real/día","Meta proteína","Falta proteína/día","Grasa real/día","Carbo real/día","Faltante kcal (días reg.)"])
dias_arr='{"Lunes","Martes","Miércoles","Jueves","Viernes","Sábado","Domingo"}'
for w in range(1,13):
    r=4+w
    cs.cell(r,1,w); cs.cell(r,2,f"=Plan_12sem!B{r}").number_format="dd/mm"
    cs.cell(r,3,f"=Plan_12sem!J{r}")
    cs.cell(r,4,f'=SUMPRODUCT(--(COUNTIFS({RG("A")},$A{r},{RG("B")},{dias_arr},{RG("H")},">0")>0))')
    cs.cell(r,5,f'=SUMIFS({RG("H")},{RG("A")},$A{r})').number_format="#,##0"
    cs.cell(r,6,f'=IF(D{r}=0,"",E{r}/D{r})').number_format="0"
    cs.cell(r,7,f'=IF(D{r}=0,"",F{r}-C{r})').number_format="+0;-0;0"
    cs.cell(r,8,f'=IF(D{r}=0,"",F{r}/C{r})').number_format="0%"
    cs.cell(r,9,f'=IF(D{r}=0,"Sin datos",IF(ABS(H{r}-1)<=0.05,"Cumple",IF(H{r}>1.05,"Exceso","Falta")))')
    cs.cell(r,10,f'=IF(D{r}=0,"",SUMIFS({RG("I")},{RG("A")},$A{r})/D{r})').number_format="0"
    cs.cell(r,11,f"=Plan_12sem!L{r}")
    cs.cell(r,12,f'=IF(D{r}=0,"",MAX(0,K{r}-J{r}))').number_format="0"
    cs.cell(r,13,f'=IF(D{r}=0,"",SUMIFS({RG("J")},{RG("A")},$A{r})/D{r})').number_format="0"
    cs.cell(r,14,f'=IF(D{r}=0,"",SUMIFS({RG("K")},{RG("A")},$A{r})/D{r})').number_format="0"
    cs.cell(r,15,f'=IF(D{r}=0,"",C{r}*D{r}-E{r})').number_format="+#,##0;-#,##0;0"
    for j in range(1,16): cs.cell(r,j).border=box
cs["A18"]="Faltante kcal: positivo = comió menos que la meta; negativo = se pasó."; cs["A18"].font=Font(italic=True)
widths(cs,[6,9,11,11,12,14,12,12,12,12,11,12,11,11,16])
rng="I5:I16"
cs.conditional_formatting.add(rng,CellIsRule(operator="equal",formula=['"Cumple"'],fill=PatternFill("solid",bgColor="C6EFCE")))
cs.conditional_formatting.add(rng,CellIsRule(operator="equal",formula=['"Exceso"'],fill=PatternFill("solid",bgColor="F8CBAD")))
cs.conditional_formatting.add(rng,CellIsRule(operator="equal",formula=['"Falta"'],fill=PatternFill("solid",bgColor="FFEB9C")))
cs.freeze_panes="B5"

# ================= PESO =================
pe=wb.create_sheet("Peso"); title(pe,"Seguimiento de peso y rendimiento","Pesar en ayunas, mismas condiciones. Rendimiento gym: 1 (bajo) a 5 (excelente).")
hdr(pe,4,["Semana","Fecha","Peso objetivo (kg)","Peso real (kg)","Dif. vs objetivo","Cambio vs sem. anterior","Estado","Rendimiento gym (1–5)","Notas"])
for w in range(0,13):
    r=5+w
    pe.cell(r,1,w)
    pe.cell(r,2,f"=Perfil!$C$13+{w}*7-1" if w>0 else "=Perfil!$C$13").number_format="dd/mm"
    pe.cell(r,3,"=Perfil!C7" if w==0 else f"=Plan_12sem!D{4+w}").number_format="0.00"
    c=pe.cell(r,4,PERFIL["peso0"] if w==0 else None); inp(c); c.number_format="0.0"
    pe.cell(r,5,f'=IF(D{r}="","",D{r}-C{r})').number_format="+0.0;-0.0;0.0"
    pe.cell(r,6,"" if w==0 else f'=IF(OR(D{r}="",D{r-1}=""),"",D{r}-D{r-1})').number_format="+0.0;-0.0;0.0"
    pe.cell(r,7,"" if w==0 else f'=IF(D{r}="","",IF(D{r}<=C{r}+0.2,"En ruta",IF(D{r}<=C{r}+0.4,"Vigilar","Revisar")))')
    inp(pe.cell(r,8)); inp(pe.cell(r,9))
    for j in range(1,10): pe.cell(r,j).border=box
pe["A19"]="Semana 0 = peso inicial (lunes 12/10). Desde la semana 1, la fecha es el domingo de cierre de esa semana."; pe["A19"].font=Font(italic=True)
widths(pe,[9,10,16,14,14,18,12,18,40])
pe.conditional_formatting.add("G5:G17",CellIsRule(operator="equal",formula=['"En ruta"'],fill=PatternFill("solid",bgColor="C6EFCE")))
pe.conditional_formatting.add("G5:G17",CellIsRule(operator="equal",formula=['"Vigilar"'],fill=PatternFill("solid",bgColor="FFEB9C")))
pe.conditional_formatting.add("G5:G17",CellIsRule(operator="equal",formula=['"Revisar"'],fill=PatternFill("solid",bgColor="F8CBAD")))
ch=LineChart(); ch.title="Peso objetivo vs real (kg)"; ch.height=8; ch.width=18
ch.add_data(Reference(pe,min_col=3,min_row=4,max_row=17),titles_from_data=True)
ch.add_data(Reference(pe,min_col=4,min_row=4,max_row=17),titles_from_data=True)
ch.set_categories(Reference(pe,min_col=1,min_row=5,max_row=17))
ch.y_axis.scaling.min=60; ch.y_axis.scaling.max=65; ch.y_axis.title="kg"; ch.x_axis.title="Semana"
ch.x_axis.delete=False; ch.y_axis.delete=False
pe.add_chart(ch,"K4")

# orden de hojas
order=["Leeme","Perfil","Plan_12sem","Alimentos","Menu_S1","Resumen_S1","Registro","Dia_Detalle","Control_Semanal","Peso","Reglas"]
wb._sheets=[wb[n] for n in order]
wb.calculation.fullCalcOnLoad=True
wb.save(OUT); print("ok",OUT)
